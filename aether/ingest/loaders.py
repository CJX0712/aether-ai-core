"""文档加载器。文本/Markdown 为零依赖；PDF/DOCX 为 lazy import 可选实现。"""

from __future__ import annotations

import os

from ..core.types import Document


def load_text(path: str) -> Document:
    with open(path, "r", encoding="utf-8") as f:
        text = f.read()
    return Document(text=text, meta={"source": path, "type": "text"}, doc_id=path)


def load_markdown(path: str) -> Document:
    d = load_text(path)
    d.meta["type"] = "markdown"
    return d


def load_pdf(path: str) -> Document:
    try:
        import pdfplumber
    except ImportError as exc:  # pragma: no cover
        raise ImportError("需要 pdfplumber：pip install pdfplumber") from exc
    pages: list[str] = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            t = page.extract_text() or ""
            if t:
                pages.append(t)
    return Document(text="\n".join(pages), meta={"source": path, "type": "pdf"}, doc_id=path)


def load_docx(path: str) -> Document:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover
        raise ImportError("需要 python-docx：pip install python-docx") from exc
    document = docx.Document(path)
    text = "\n".join(p.text for p in document.paragraphs if p.text)
    return Document(text=text, meta={"source": path, "type": "docx"}, doc_id=path)


def load_auto(path: str) -> Document:
    ext = os.path.splitext(path)[1].lower()
    if ext in (".md", ".markdown"):
        return load_markdown(path)
    if ext == ".pdf":
        return load_pdf(path)
    if ext == ".docx":
        return load_docx(path)
    return load_text(path)


__all__ = ["load_text", "load_markdown", "load_pdf", "load_docx", "load_auto"]
